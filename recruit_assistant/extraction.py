import hashlib
import re
from io import BytesIO
from pathlib import Path
from urllib.parse import urlparse, urlunparse

from docx import Document
from pypdf import PdfReader
from pypdf.errors import PyPdfError

URL_PATTERN = re.compile(
    r"(?:https?://)?(?:www\.)?(?:github\.com|linkedin\.com/in)/[A-Za-z0-9_.~/%?#@!$&'()*+,;=:-]+",
    re.IGNORECASE,
)


def extract_resume_text(content: bytes, extension: str) -> str:
    if extension == ".txt":
        text = content.decode("utf-8", errors="replace")
    elif extension == ".pdf":
        try:
            reader = PdfReader(BytesIO(content))
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        except PyPdfError as exc:
            raise ValueError(
                "Invalid or corrupted PDF. Re-export the original document as a "
                "PDF and upload it again."
            ) from exc
    elif extension == ".docx":
        document = Document(BytesIO(content))
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        table_cells = [
            cell.text
            for table in document.tables
            for row in table.rows
            for cell in row.cells
        ]
        text = "\n".join(paragraphs + table_cells)
    elif extension == ".doc":
        raise ValueError(
            "Legacy .doc text extraction is not supported locally. Convert it to "
            ".docx, .pdf, or .txt."
        )
    else:
        raise ValueError(f"Unsupported resume extension: {extension}")

    normalized = normalize_text(text)
    if not normalized:
        raise ValueError(
            "No text could be extracted. The document may be image-only and require OCR."
        )
    return normalized


def normalize_text(text: str) -> str:
    text = text.replace("\x00", "")
    lines = [" ".join(line.split()) for line in text.splitlines()]
    return "\n".join(line for line in lines if line).strip()


def extract_professional_links(text: str) -> list[str]:
    links: set[str] = set()
    for match in URL_PATTERN.findall(text):
        raw = match.rstrip(".,);]}>")
        if not raw.lower().startswith(("http://", "https://")):
            raw = f"https://{raw}"
        parsed = urlparse(raw)
        host = parsed.netloc.lower().removeprefix("www.")
        path = parsed.path.rstrip("/")
        normalized = urlunparse(("https", host, path, "", "", ""))
        links.add(normalized)
    return sorted(links)


def extract_embedded_professional_links(
    content: bytes,
    extension: str,
) -> list[str]:
    targets: list[str] = []
    try:
        if extension == ".pdf":
            reader = PdfReader(BytesIO(content))
            for page in reader.pages:
                for annotation_ref in page.get("/Annots", []):
                    annotation = annotation_ref.get_object()
                    action = annotation.get("/A")
                    if action and action.get("/URI"):
                        targets.append(str(action["/URI"]))
        elif extension == ".docx":
            document = Document(BytesIO(content))
            targets.extend(
                str(relationship.target_ref)
                for relationship in document.part.rels.values()
                if relationship.is_external
            )
    except (PyPdfError, OSError, ValueError):
        return []
    return extract_professional_links("\n".join(targets))


def text_hash(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def write_extracted_text(job_dir: Path, resume_id: str, text: str) -> str:
    relative_path = Path("extracted") / f"{resume_id}.txt"
    path = job_dir / relative_path
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding="utf-8")
    return str(relative_path)
